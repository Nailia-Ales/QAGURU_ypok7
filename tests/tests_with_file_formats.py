from zipfile import ZipFile

from openpyxl.reader.excel import load_workbook
from pypdf import PdfReader

from tests.conftest import CURRENT_PROJECT_PATH, ARCHIVE_FILE_PATH


# from tests.conftest import


def test_pdf_reader():
    with ZipFile(ARCHIVE_FILE_PATH, "r") as zip_file:
        with zip_file.open("pdf_file.pdf") as pdf_file:
            reader = PdfReader(pdf_file)
            assert "Создание  опроса" in reader.pages[0].extract_text()


def test_xlsx_reader():
    with ZipFile(ARCHIVE_FILE_PATH, "r") as zip_file:
        with zip_file.open("xlsx_file.xlsx") as xlsx_file:
            reader = load_workbook(xlsx_file)
            assert "погода" in reader.active.cell(2, 1).value


def test_csv_reader():
    with ZipFile(ARCHIVE_FILE_PATH, "r") as zip_file:
        with zip_file.open("csv_file.csv") as csv_file:
            reader = csv_file.read().decode("utf-8")
            assert "OU005" in reader